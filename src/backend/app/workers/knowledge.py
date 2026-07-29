import io
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.logging import get_logger
from app.core.security import now
from app.core.text import truncate
from app.db.models.enums import FileStatus
from app.db.models.job import Job
from app.db.models.knowledge import KnowledgeBase, KnowledgeFile, KnowledgeSource
from app.integrations.maas import MaasError, maas
from app.integrations.search import search_client
from app.services.storage import storage
from app.workers.base import JobCancelled, check_cancel, claim, fail, finish, step

log = get_logger(__name__)

SUMMARY_SYSTEM = """\
Bạn là trợ lý quản lý tri thức. Đọc nội dung tài liệu và viết tóm tắt 2-3 câu
bằng tiếng Việt, nêu rõ tài liệu này dùng để làm gì và chứa thông tin gì.
Chỉ trả về đoạn tóm tắt, không thêm lời dẫn.
"""


async def index_document(ctx: dict[str, Any], file_id: str) -> None:
    from app.db.session import SessionLocal

    async with SessionLocal() as session:
        row = await session.get(KnowledgeFile, file_id)
        if row is None:
            return

        row.status = FileStatus.INDEXING
        await session.commit()

        try:
            data = await storage.get(settings.bucket_knowledge, row.object_key)
            text, pages = await _extract(row, data)

            row.extracted_text = text or None
            row.page_count = pages
            row.status = FileStatus.INDEXED
            row.indexed_at = now()
            row.index_error = None

            if text and maas.configured:
                try:
                    completion = await maas.run(
                        "kb.summary",
                        SUMMARY_SYSTEM,
                        f"TÊN TỆP: {row.name}\n\nNỘI DUNG:\n{truncate(text, 12000)}",
                    )
                    row.summary = truncate(completion.text.strip(), 800)
                except MaasError as exc:
                    log.warning("summary_failed", file=file_id, error=str(exc))

            await session.commit()
            log.info("document_indexed", file=file_id, chars=len(text or ""), pages=pages)

        except Exception as exc:
            log.warning("index_failed", file=file_id, error=str(exc))
            row.status = FileStatus.FAILED
            row.index_error = str(exc)[:500]
            await session.commit()


async def _extract(row: KnowledgeFile, data: bytes) -> tuple[str, int | None]:
    ext = (row.ext or "").lower()

    if ext == "pdf":
        return await _extract_pdf(data)
    if ext == "docx":
        return _extract_docx(data), None
    if ext == "xlsx":
        return _extract_xlsx(data), None
    if ext == "pptx":
        return _extract_pptx(data)
    return "", None


async def _extract_pdf(data: bytes) -> tuple[str, int | None]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = len(reader.pages)
    chunks: list[str] = []
    for page in reader.pages:
        try:
            chunks.append(page.extract_text() or "")
        except Exception:
            chunks.append("")

    text = "\n".join(chunks).strip()

    if len(text) < 200 and maas.configured:
        try:
            completion = await maas.look(
                "vision.ocr",
                "Đọc toàn bộ chữ trong tài liệu này và trả về dưới dạng văn bản thuần. "
                "Giữ nguyên thứ tự, không thêm bình luận.",
                [data],
                mime_type="application/pdf",
            )
            vision_text = completion.text.strip()
            if len(vision_text) > len(text):
                return vision_text, pages
        except MaasError as exc:
            log.warning("vision_extract_failed", error=str(exc))

    return text, pages


def _extract_docx(data: bytes) -> str:
    import re
    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
    except (KeyError, zipfile.BadZipFile) as exc:
        raise ValueError(f"DOCX không đọc được: {exc}") from exc

    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"<w:tab[^>]*/>", "\t", xml)
    text = re.sub(r"<[^>]+>", "", xml)
    return _unescape(text).strip()


def _extract_xlsx(data: bytes) -> str:
    import re
    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
            shared: list[str] = []
            if "xl/sharedStrings.xml" in names:
                raw = archive.read("xl/sharedStrings.xml").decode("utf-8", errors="replace")
                shared = [
                    _unescape(re.sub(r"<[^>]+>", "", si))
                    for si in re.findall(r"<si>(.*?)</si>", raw, re.S)
                ]

            lines: list[str] = []
            for name in sorted(n for n in names if n.startswith("xl/worksheets/sheet")):
                sheet = archive.read(name).decode("utf-8", errors="replace")
                for row in re.findall(r"<row[^>]*>(.*?)</row>", sheet, re.S):
                    cells: list[str] = []
                    for cell in re.findall(r"<c[^>]*?(?:\st=\"(\w)\")?[^>]*>(.*?)</c>", row, re.S):
                        kind, body = cell
                        value = _unescape(re.sub(r"<[^>]+>", "", body)).strip()

                        if kind == "s" and value.isdigit() and int(value) < len(shared):
                            value = shared[int(value)]
                        if value:
                            cells.append(value)
                    if cells:
                        lines.append("\t".join(cells))
            return "\n".join(lines).strip()
    except zipfile.BadZipFile as exc:
        raise ValueError(f"XLSX không đọc được: {exc}") from exc


def _extract_pptx(data: bytes) -> tuple[str, int]:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(data))
    chunks: list[str] = []
    count = 0
    for index, slide in enumerate(deck.slides, start=1):
        count = index
        parts = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if parts:
            chunks.append(f"[Slide {index}]\n" + "\n".join(parts))
    return "\n\n".join(chunks).strip(), count


def _unescape(text: str) -> str:
    return (
        text.replace("&amp;", "&")
        .replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&quot;", '"')
        .replace("&apos;", "'")
    )


async def run_kb_sync(ctx: dict[str, Any], job_id: str) -> None:
    from app.db.session import SessionLocal

    async with SessionLocal() as session:
        job = await claim(session, job_id)
        if job is None:
            return
        try:
            await _sync(session, job)
        except JobCancelled:
            return
        except Exception as exc:
            log.exception("kb_sync_crashed", job=job_id, error=str(exc))
            await fail(session, job, "KB_SYNC_FAILED", f"Không đồng bộ được: {exc}"[:280])


async def _sync(session: AsyncSession, job: Job) -> None:
    kb = await session.get(KnowledgeBase, job.kb_id) if job.kb_id else None
    if kb is None:
        await fail(session, job, "KB_SYNC_FAILED", "Không tìm thấy kho tri thức.")
        return

    sources = list(
        (await session.execute(select(KnowledgeSource).where(KnowledgeSource.kb_id == kb.id)))
        .scalars()
        .all()
    )
    if not sources:
        await finish(session, job, {"syncedCount": 0, "message": "Kho chưa có nguồn nào."})
        return

    synced = 0
    failed = 0
    for index, source in enumerate(sources, start=1):
        await check_cancel(session, job)
        await step(
            session,
            job,
            "sync",
            f"Đang đồng bộ {source.url}",
            index / max(1, len(sources)) * 0.9,
        )

        outcome = await search_client.search(f"site:{source.url}", limit=5)
        if outcome.ok and outcome.hits:
            source.status = "ok"
            source.doc_count = len(outcome.hits)
            source.last_synced_at = now()
            source.last_error = None
            source.meta = {"pages": [{"url": h.url, "title": h.title} for h in outcome.hits[:10]]}
            synced += 1
        else:
            source.status = "failed"
            source.last_error = outcome.reason or "no_result"
            failed += 1
        await session.commit()

    await finish(
        session, job, {"syncedCount": synced, "failedCount": failed, "total": len(sources)}
    )
